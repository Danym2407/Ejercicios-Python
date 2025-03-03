import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo

# Definimos las materias y sus horarios
materias = [
    ('LABORATORIO DE FUNDAMENTOS DE PROGRAMACIÓN M', 'Martes', '9 - 11 Hrs.'),
    ('LABORATORIO DE FUNDAMENTOS DE PROGRAMACIÓN L', 'Lunes', '18 - 20 Hrs.'),
    ('LABORATORIO DE FUNDAMENTOS DE PROGRAMACIÓN V', 'Viernes', '7 - 9 Hrs.'),
    ('LABORATORIO DE FUNDAMENTOS DE PROGRAMACIÓN V', 'Viernes', '9 - 11 Hrs.'),
    ('LABORATORIO DE INGENIERÍA DE CIRCUITOS I HERNANDEZ MONROY JESSICA M', 'Miércoles', '7 - 9 Hrs.'),
    ('LABORATORIO DE INGENIERÍA DE CIRCUITOS I HERNANDEZ MONROY JESSICA V', 'Viernes', '7 - 9 Hrs.'),
    ('LABORATORIO DE INGENIERÍA DE CIRCUITOS I LOPEZ GONZALEZ ALEXANDRO M', 'Martes', '16 - 18 Hrs.'),
    ('LABORATORIO DE INGENIERÍA DE CIRCUITOS I VANDER MERSCH HUERTA ROMO BERNARD ROELAND L', 'Lunes', '11 - 13 Hrs.'),
    ('LABORATORIO DE PROGRAMACIÓN APLICADA V', 'Viernes', '9 - 11 Hrs.'),
    ('LABORATORIO DE PROGRAMACIÓN APLICADA V', 'Viernes', '11 - 13 Hrs.'),
    ('LABORATORIO DE PROGRAMACIÓN APLICADA V', 'Viernes', '11 - 13 Hrs.'),
    ('LABORATORIO DE PROGRAMACIÓN APLICADA M', 'Miércoles', '11 - 13 Hrs.'),
    ('LAB. DE DISEÑO DE SISTEMAS DIGITALES M', 'Miércoles', '7 - 11 Hrs.'),
    ('LAB. DE DISEÑO DE SISTEMAS DIGITALES M', 'Miércoles', '7 - 11 Hrs.'),
    ('LABORATORIO DE INGENIERÍA DE CIRCUITOS II WITEMBERG WUDKA NATHAN M', 'Martes', '16 - 18 Hrs.'),
    ('LABORATORIO DE INGENIERÍA DE CIRCUITOS II WITEMBERG WUDKA NATHAN M', 'Jueves', '16 - 18 Hrs.'),
    ('LABORATORIO DE INGENIERÍA DE CIRCUITOS II MARTINEZ CERVANTES LUIS MIGUEL L', 'Lunes', '9 - 11 Hrs.'),
    ('LABORATORIO DE INGENIERÍA DE CIRCUITOS II MARTINEZ CERVANTES LUIS MIGUEL L', 'Miércoles', '9 - 11 Hrs.'),
    ('LAB. DE ING', 'Miércoles', '7 - 9 Hrs.'),
    ('LAB. DE ING', 'Viernes', '7 - 9 Hrs.'),
    ('LABORATORIO DE SISTEMAS DE BASES DE DATOS L', 'Lunes', '7 - 9 Hrs.'),
    ('LABORATORIO DE SISTEMAS DE BASES DE DATOS M', 'Martes', '20 - 21 Hrs.'),
    ('LABORATORIO DE SISTEMAS DE BASES DE DATOS M', 'Jueves', '20 - 21 Hrs.'),
    ('LAB. DE PROGRAMACIÓN ORIENTADA A OBJETOS L', 'Lunes', '16 - 18 Hrs.'),
    ('LAB. DE PROGRAMACIÓN ORIENTADA A OBJETOS L', 'Lunes', '16 - 18 Hrs.'),
    ('LABORATORIO DE SENSORES Y ACTUADORES L', 'Lunes', '9 - 11 Hrs.'),
    ('LABORATORIO DE SISTEMAS DE COMUNICACIONES M', 'Martes', '9 - 10 Hrs.'),
    ('LABORATORIO DE SISTEMAS DE COMUNICACIONES M', 'Jueves', '9 - 10 Hrs.'),
    ('LAB. DE FUNDAMENTOS DE REDES DIGITALES V', 'Viernes', '11 - 13 Hrs.'),
    ('LABORATORIO DE INGENIERÍA DE POTENCIA MARTINEZ CERVANTES LUIS MIGUEL L', 'Lunes', '9 - 11 Hrs.'),
    ('LABORATORIO DE INGENIERÍA DE POTENCIA MARTINEZ CERVANTES LUIS MIGUEL L', 'Miércoles', '9 - 11 Hrs.'),
    ('LAB. DE ARQUITECTURA DE INFORMACIÓN WEB L', 'Lunes', '7 - 9 Hrs.'),
    ('LAB. DE REDES INALÁMBRICAS Y MÓVILES VEGA LIMON ANDRES S', 'Sábado', '8 - 10 Hrs.'),
    ('LABORATORIO DE APLICACIONES MÓVILES V', 'Viernes', '16 - 18 Hrs.'),
    ('LAB. DE AUTOMATIZACIÓN DE SIST', 'Miércoles', '9 - 11 Hrs.'),
    ('LABORATORIO DE PROCESAMIENTO DE IMÁGENES V', 'Viernes', '11 - 13 Hrs.'),
    ('LABORATORIO DE PROCESAMIENTO DE SEÑALES M', 'Martes', '7 - 8 Hrs.'),
    ('LABORATORIO DE PROCESAMIENTO DE SEÑALES M', 'Jueves', '7 - 8 Hrs.'),
    ('LABORATORIO DE SISTEMAS ANALÓGICOS WITEMBERG WUDKA NATHAN M', 'Martes', '16 - 18 Hrs.'),
    ('LABORATORIO DE SISTEMAS ANALÓGICOS WITEMBERG WUDKA NATHAN M', 'Jueves', '16 - 18 Hrs.'),
    ('LABORATORIO DE SISTEMAS ANALÓGICOS MARTINEZ CERVANTES LUIS MIGUEL L', 'Lunes', '9 - 11 Hrs.'),
    ('LABORATORIO DE SISTEMAS ANALÓGICOS MARTINEZ CERVANTES LUIS MIGUEL L', 'Miércoles', '9 - 11 Hrs.'),
    ('LAB. DE INGENIERÍA DE CIRCUITOS M', 'Miércoles', '7 - 9 Hrs.'),
    ('LAB. DE INGENIERÍA DE CIRCUITOS V', 'Viernes', '7 - 9 Hrs.'),
    ('LAB. DE INGENIERÍA DE CIRCUITOS M', 'Martes', '16 - 18 Hrs.'),
    ('LAB. DE INGENIERÍA DE CIRCUITOS L', 'Lunes', '11 - 13 Hrs.'),
    ('LAB. DE ANALÓGICOS Y DE POTENCIA M', 'Martes', '16 - 18 Hrs.'),
    ('LAB. DE ANALÓGICOS Y DE POTENCIA M', 'Jueves', '16 - 18 Hrs.'),
    ('LAB. DE ANALÓGICOS Y DE POTENCIA L', 'Lunes', '9 - 11 Hrs.'),
    ('LAB. DE ANALÓGICOS Y DE POTENCIA L', 'Miércoles', '9 - 11 Hrs.'),
    ('LAB. DE SISTEMAS DIGITALES VANDER MERSCH HUERTA ROMO BERNARD ROELAND M', 'Miércoles', '7 - 11 Hrs.'),
    ('LAB. DE SISTEMAS DIGITALES DOUNCE PEREZ TAGLE CARLOS ARTURO M', 'Miércoles', '7 - 11 Hrs.'),
    ('LAB. DE SISTEMAS EMBEBIDOS J', 'Jueves', '13 - 15 Hrs.'),
    ('LAB. DE SISTEMAS EMBEBIDOS M', 'Martes', '13 - 15 Hrs.'),
    ('LAB. DE INGENIERÍA DE SOFTWARE V', 'Viernes', '9 - 11 Hrs.'),
    ('LAB. DE ARQUITECTURA DE SISTEMAS EMBEBIDOS J', 'Jueves', '16 - 18 Hrs.'),
    ('LAB. DE APLICACIONES DE REDES V', 'Viernes', '18 - 20 Hrs.'),
    ('LAB. DE SEGURIDAD DE LA INFORMACIÓN M', 'Martes', '17 - 18 Hrs.'),
    ('LAB. DE SEGURIDAD DE LA INFORMACIÓN M', 'Jueves', '17 - 18 Hrs.'),
    ('LAB. DE COM INALÁMBRICAS Y ÓPTICAS M', 'Martes', '16 - 18 Hrs.'),
    ('LAB. DE REDES DE TELECOMUNICACIONES M', 'Miércoles', '18 - 20 Hrs.'),
    ('LAB. DE REDES INALÁMBRICAS S', 'Sábado', '8 - 10 Hrs.'),
    ('LAB DE AUTOMATIZACIÓN POR LÓGICA PROGRAMA', 'Miércoles', '7 - 9 Hrs.'),
    ('LAB DE AUTOMATIZACIÓN DE SISTEMAS PRODUCT', 'Lunes', '9 - 11 Hrs.'),
    ('LAB DE CONTROL J', 'Jueves', '9 - 11 Hrs.'),
    ('LAB DE CONTROL AVANZADO V', 'Viernes', '7 - 9 Hrs.'),
    ('LAB DE ROBÓTICA MÓVIL J', 'Jueves', '7 - 9 Hrs.'),
    ('LAB DE INTEGRACIÓN MECATRÓNICA V', 'Viernes', '13 - 15 Hrs.'),
    ('LAB DE AERONAVES NO TRIPULADAS J', 'Jueves', '9 - 11 Hrs.'),
]


# Crear ventana
root = tk.Tk()
root.title("Horario Visual")

# Crear un canvas y un scrollbar
canvas = tk.Canvas(root)
scroll_y = tk.Scrollbar(root, orient="vertical", command=canvas.yview)
canvas.configure(yscrollcommand=scroll_y.set)

# Crear un frame para el contenido
frame = tk.Frame(canvas)

# Crear la estructura de la tabla
days = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado']
times = [f"{i}:00" for i in range(7, 23)]  # Horarios de 7am a 10pm (en formato de 24 horas)

# Crear un diccionario para almacenar las celdas
schedule = {day: {time: "" for time in times} for day in days}

# Llenar el horario con las materias
for materia, dia, horario in materias:
    dia_list = dia.split(" y ")  # Si el día tiene 'y', los divide en dos
    horas = horario.split(" - ")

    # Manejo de horarios
    if len(horas) == 2:  # Caso en que el horario es un rango
        start_hour = int(horas[0].split()[0])
        end_hour = int(horas[1].split()[0])
    else:  # Caso en que el horario es una sola hora (e.g., "16 Hrs.")
        start_hour = int(horas[0].split()[0])
        end_hour = start_hour + 1  # Asumimos que dura una hora

    # Recorrer los días y asignar la materia a la hora correspondiente
    for day in dia_list:
        for time in range(start_hour, end_hour):
            # Concatenar las materias, sin sobrescribir lo existente
            if schedule[day][f"{time}:00"] == "":
                schedule[day][f"{time}:00"] = materia
            else:
                schedule[day][f"{time}:00"] += f"\n{materia}"

# Crear los encabezados de la tabla
for i, day in enumerate(days):
    label = tk.Label(frame, text=day, relief="solid", width=20, height=2)
    label.grid(row=0, column=i+1, padx=5, pady=5)

# Crear las celdas para cada hora y día
for i, time in enumerate(times):
    label = tk.Label(frame, text=time, relief="solid", width=20, height=2)
    label.grid(row=i+1, column=0, padx=5, pady=5)

    for j, day in enumerate(days):
        subject = schedule[day][time]
        label = tk.Label(frame, text=subject, relief="solid", width=20, height=4, anchor="nw", justify="left")
        label.grid(row=i+1, column=j+1, padx=5, pady=5)

# Agregar el frame al canvas y la scrollbar al canvas
canvas.create_window((0, 0), window=frame, anchor="nw")
scroll_y.pack(side="right", fill="y")
canvas.pack(side="left", fill="both", expand=True)

# Configurar el scroll de acuerdo al tamaño del contenido
frame.update_idletasks()  # Necesario para calcular el tamaño total del contenido
canvas.config(scrollregion=canvas.bbox("all"))

# Guardar el horario en un archivo Excel
wb = Workbook()
ws = wb.active
ws.title = "Horario"

# Crear encabezado en el archivo Excel
ws.append(["Horario/Día"] + days)

# Escribir las celdas de tiempo y materias en el archivo Excel
for i, time in enumerate(times):
    row = [time] + [schedule[day][time] for day in days]
    ws.append(row)

# Estilizar la tabla en el archivo Excel
for row in ws.iter_rows(min_row=1, max_row=len(times) + 1, min_col=1, max_col=len(days) + 1):
    for cell in row:
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if cell.row == 1:
            cell.font = Font(bold=True)

# Crear la tabla en Excel para mejor visualización
table = Table(displayName="Horario", ref=f"A1:{chr(65 + len(days))}{len(times) + 1}")
ws.add_table(table)

# Guardar el archivo Excel
wb.save("horario.xlsx")

# Mostrar mensaje de éxito
messagebox.showinfo("Éxito", "El archivo Excel ha sido creado con éxito.")

# Iniciar la aplicación
root.mainloop()
